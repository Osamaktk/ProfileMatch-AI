from __future__ import annotations

import json
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


NAVY = "173B53"
BLUE = "147CC1"
PALE_BLUE = "EAF4FB"
GREEN = "DDF3E8"
GREEN_TEXT = "176B4D"
AMBER = "FFF0D8"
AMBER_TEXT = "9A5B00"
RED = "FCE2E0"
RED_TEXT = "A33A32"
GRAY = "EEF2F5"
GRAY_TEXT = "526779"
WHITE = "FFFFFF"
BORDER = Border(bottom=Side(style="thin", color="D8E2EA"))


CONTACT_COLUMNS = [
    "Row",
    "Name",
    "Organization",
    "Original Title",
    "Email",
    "Status",
    "Flag",
    "LinkedIn Profile",
    "Verified Role",
    "Role & Background",
    "City",
    "State",
    "Department Bucket",
    "Closest Sub-Function",
    "Local Context",
    "Flag Reason",
    "Primary Evidence",
    "Evidence Count",
]


def export_enrichment_workbook(
    contacts: list[dict[str, Any]],
    stored_results: list[dict[str, Any]],
    original_headers: list[str],
    path: Path,
) -> Path:
    """Create a review-friendly workbook while preserving the uploaded source data."""

    result_map = {int(item["row_id"]): item for item in stored_results}
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    contacts_sheet = workbook.create_sheet("Enriched Contacts")
    evidence_sheet = workbook.create_sheet("Evidence")
    source_sheet = workbook.create_sheet("Source Data")

    _write_summary(summary_sheet, contacts, result_map)
    _write_contacts(contacts_sheet, contacts, result_map)
    _write_evidence(evidence_sheet, contacts, result_map)
    _write_source_data(source_sheet, contacts, result_map, original_headers)

    workbook.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(temporary)
        workbook.close()
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
    return path


def _write_summary(
    sheet: Any,
    contacts: list[dict[str, Any]],
    result_map: dict[int, dict[str, Any]],
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F2")
    title = sheet["A1"]
    title.value = "CUSTOMER EMPOWER · CONTACT ENRICHMENT REPORT"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(color=WHITE, bold=True, size=18)
    title.alignment = Alignment(vertical="center", horizontal="left")
    for row in sheet["A1:F2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY)

    generated = datetime.now(timezone.utc).strftime("Generated %Y-%m-%d %H:%M UTC")
    sheet["A3"] = generated
    sheet["A3"].font = Font(color=GRAY_TEXT, italic=True, size=10)
    sheet.merge_cells("A3:F3")

    statuses = Counter(
        str(result_map.get(int(contact["row_id"]), {}).get("status") or "pending")
        for contact in contacts
    )
    metrics = [
        ("Total contacts", len(contacts), NAVY),
        ("Completed", statuses["completed"], GREEN_TEXT),
        ("Review required", statuses["manual_review"], AMBER_TEXT),
        ("Pending", statuses["pending"] + statuses["processing"], GRAY_TEXT),
        ("Failed", statuses["failed"], RED_TEXT),
    ]
    sheet["A5"] = "PROCESSING OVERVIEW"
    sheet["A5"].font = Font(color=NAVY, bold=True, size=12)
    for index, (label, value, color) in enumerate(metrics, start=6):
        sheet.cell(index, 1, label).font = Font(color=GRAY_TEXT, bold=True)
        sheet.cell(index, 2, value).font = Font(color=color, bold=True, size=14)
        sheet.cell(index, 1).border = BORDER
        sheet.cell(index, 2).border = BORDER

    buckets = Counter()
    for item in result_map.values():
        result = item.get("result") or {}
        bucket = result.get("department_bucket")
        if bucket:
            buckets[str(bucket)] += 1
    sheet["D5"] = "DEPARTMENT CLASSIFICATION"
    sheet["D5"].font = Font(color=NAVY, bold=True, size=12)
    sheet["D6"] = "Department"
    sheet["E6"] = "Contacts"
    _style_header_row(sheet, 6, 4, 5)
    department_row = 7
    for department, count in sorted(buckets.items(), key=lambda item: (-item[1], item[0])):
        sheet.cell(department_row, 4, department)
        sheet.cell(department_row, 5, count)
        sheet.cell(department_row, 4).border = BORDER
        sheet.cell(department_row, 5).border = BORDER
        department_row += 1
    if not buckets:
        sheet.cell(department_row, 4, "No classifications yet")
        sheet.cell(department_row, 4).font = Font(color=GRAY_TEXT, italic=True)

    note_row = max(13, department_row + 2)
    sheet.cell(note_row, 1, "HOW TO USE THIS WORKBOOK")
    sheet.cell(note_row, 1).font = Font(color=NAVY, bold=True, size=12)
    notes = [
        "Enriched Contacts contains the focused review output and clickable profile/evidence links.",
        "Evidence stores one public source per row; it replaces long lists of URLs inside contact cells.",
        "Source Data preserves every column and value from the uploaded file.",
        "Review Required means identity, role, department, or profile completeness needs a person to check it.",
    ]
    for offset, note in enumerate(notes, start=1):
        cell = sheet.cell(note_row + offset, 1, f"• {note}")
        sheet.merge_cells(start_row=note_row + offset, start_column=1, end_row=note_row + offset, end_column=6)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(color=GRAY_TEXT, size=10)

    for column, width in {"A": 28, "B": 14, "C": 4, "D": 34, "E": 14, "F": 12}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"


def _write_contacts(
    sheet: Any,
    contacts: list[dict[str, Any]],
    result_map: dict[int, dict[str, Any]],
) -> None:
    for column, header in enumerate(CONTACT_COLUMNS, start=1):
        sheet.cell(1, column, header)
    _style_header_row(sheet, 1, 1, len(CONTACT_COLUMNS))

    for output_row, contact in enumerate(contacts, start=2):
        stored = result_map.get(int(contact["row_id"]), {})
        result = stored.get("result") or {}
        original = dict(contact.get("original_data") or {})
        evidence = _result_evidence(result)
        primary = max(evidence, key=_evidence_score) if evidence else None
        values = [
            contact.get("source_row") or contact.get("row_id"),
            contact.get("name") or _pick(original, "Name", "Contact Name"),
            contact.get("organization") or _pick(original, "Organization", "Company"),
            contact.get("title") or _pick(original, "Title", "Job Title"),
            contact.get("email") or _pick(original, "Email", "Email Address"),
            _status_label(stored.get("status")),
            _flag_label(stored, result),
            result.get("linkedin_url"),
            result.get("role"),
            result.get("background_summary"),
            result.get("city"),
            result.get("state"),
            result.get("department_bucket"),
            result.get("sub_function"),
            result.get("local_context"),
            result.get("review_reason") or stored.get("error"),
            _evidence_label(primary) if primary else None,
            len(evidence),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(output_row, column, _safe_value(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        _set_hyperlink(sheet.cell(output_row, 8), result.get("linkedin_url"))
        if primary:
            _set_hyperlink(sheet.cell(output_row, 17), primary.get("source_url"))
        _style_status(sheet.cell(output_row, 6), stored.get("status"))
        _style_flag(sheet.cell(output_row, 7), _flag_label(stored, result))
        sheet.row_dimensions[output_row].height = 62 if result else 30

    _add_table(sheet, "EnrichedContactsTable", len(contacts) + 1, len(CONTACT_COLUMNS))
    widths = {
        1: 8,
        2: 24,
        3: 30,
        4: 28,
        5: 30,
        6: 18,
        7: 12,
        8: 36,
        9: 30,
        10: 58,
        11: 20,
        12: 10,
        13: 24,
        14: 28,
        15: 56,
        16: 52,
        17: 42,
        18: 14,
    }
    _finish_data_sheet(sheet, widths, freeze="A2", landscape=True)


def _write_evidence(
    sheet: Any,
    contacts: list[dict[str, Any]],
    result_map: dict[int, dict[str, Any]],
) -> None:
    headers = [
        "Row",
        "Contact",
        "Organization",
        "Source Type",
        "Source Title",
        "Evidence Snippet",
        "Source URL",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(1, column, header)
    _style_header_row(sheet, 1, 1, len(headers))

    output_row = 2
    for contact in contacts:
        stored = result_map.get(int(contact["row_id"]), {})
        result = stored.get("result") or {}
        evidence = sorted(_result_evidence(result), key=_evidence_score, reverse=True)
        for item in evidence:
            values = [
                contact.get("source_row") or contact.get("row_id"),
                contact.get("name"),
                contact.get("organization"),
                _friendly_source_type(item.get("source_type") or item.get("evidence_type")),
                _evidence_label(item),
                item.get("extracted_text"),
                item.get("source_url"),
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(output_row, column, _safe_value(value))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = BORDER
            _set_hyperlink(sheet.cell(output_row, 5), item.get("source_url"))
            _set_hyperlink(sheet.cell(output_row, 7), item.get("source_url"))
            sheet.row_dimensions[output_row].height = 48
            output_row += 1

    if output_row == 2:
        sheet.cell(2, 1, "No evidence has been collected yet.")
    else:
        _add_table(sheet, "EvidenceTable", output_row - 1, len(headers))
    _finish_data_sheet(
        sheet,
        {1: 8, 2: 24, 3: 30, 4: 22, 5: 44, 6: 72, 7: 62},
        freeze="A2",
        landscape=True,
    )


def _write_source_data(
    sheet: Any,
    contacts: list[dict[str, Any]],
    result_map: dict[int, dict[str, Any]],
    original_headers: list[str],
) -> None:
    headers = _unique_headers([str(header or "Column") for header in original_headers])
    output_headers = [*headers, "Enrichment Status", "Flag"]
    for column, header in enumerate(output_headers, start=1):
        sheet.cell(1, column, header)
    _style_header_row(sheet, 1, 1, len(output_headers))

    for output_row, contact in enumerate(contacts, start=2):
        original = dict(contact.get("original_data") or {})
        stored = result_map.get(int(contact["row_id"]), {})
        for column, header in enumerate(headers, start=1):
            value = original.get(original_headers[column - 1], "")
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            cell = sheet.cell(output_row, column, _safe_value(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        status_column = len(headers) + 1
        status_cell = sheet.cell(output_row, status_column, _status_label(stored.get("status")))
        status_cell.alignment = Alignment(vertical="top")
        status_cell.border = BORDER
        _style_status(status_cell, stored.get("status"))
        result = stored.get("result") or {}
        flag_value = _flag_label(stored, result)
        flag_cell = sheet.cell(output_row, status_column + 1, flag_value)
        flag_cell.border = BORDER
        _style_flag(flag_cell, flag_value)

    _add_table(sheet, "SourceDataTable", len(contacts) + 1, len(output_headers))
    widths = {index: 20 for index in range(1, len(output_headers) + 1)}
    for index, header in enumerate(output_headers, start=1):
        if header.casefold() in {"name", "organization", "company"}:
            widths[index] = 28
        elif "email" in header.casefold():
            widths[index] = 30
        elif header == "Enrichment Status":
            widths[index] = 18
        elif header == "Flag":
            widths[index] = 12
    _finish_data_sheet(sheet, widths, freeze="A2", landscape=True)


def _style_header_row(sheet: Any, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = sheet.cell(row, column)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 34


def _finish_data_sheet(
    sheet: Any,
    widths: dict[int, float],
    *,
    freeze: str,
    landscape: bool,
) -> None:
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = freeze
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.sheet_properties.outlinePr.summaryBelow = True


def _add_table(sheet: Any, name: str, max_row: int, max_column: int) -> None:
    if max_row < 2 or max_column < 1:
        return
    reference = f"A1:{get_column_letter(max_column)}{max_row}"
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _style_status(cell: Any, status: str | None) -> None:
    style = {
        "completed": (GREEN, GREEN_TEXT),
        "manual_review": (AMBER, AMBER_TEXT),
        "failed": (RED, RED_TEXT),
        "processing": (PALE_BLUE, BLUE),
        "pending": (GRAY, GRAY_TEXT),
    }.get(str(status), (GRAY, GRAY_TEXT))
    cell.fill = PatternFill("solid", fgColor=style[0])
    cell.font = Font(color=style[1], bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_flag(cell: Any, value: str) -> None:
    if value == "YES":
        colors = (RED, RED_TEXT)
    elif value == "NO":
        colors = (GREEN, GREEN_TEXT)
    else:
        colors = (GRAY, GRAY_TEXT)
    cell.fill = PatternFill("solid", fgColor=colors[0])
    cell.font = Font(color=colors[1], bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _flag_label(stored: dict[str, Any], result: dict[str, Any]) -> str:
    status = str(stored.get("status") or "pending")
    if status in {"manual_review", "failed"} or result.get("manual_review") is True:
        return "YES"
    if status == "completed" and result:
        return "NO"
    return ""


def _set_hyperlink(cell: Any, value: Any) -> None:
    text = str(value or "").strip()
    if text.startswith(("http://", "https://")):
        cell.hyperlink = text
        cell.font = Font(color="0563C1", underline="single")


def _result_evidence(result: dict[str, Any]) -> list[dict[str, Any]]:
    evidence = [item for item in (result.get("evidence") or []) if isinstance(item, dict)]
    if evidence:
        return evidence
    return [
        {
            "source_url": url,
            "source_title": None,
            "extracted_text": None,
            "source_type": "public_source",
        }
        for url in (result.get("sources") or [])
        if isinstance(url, str) and url.startswith(("http://", "https://"))
    ]


def _evidence_score(item: dict[str, Any]) -> int:
    url = str(item.get("source_url") or "")
    parsed = urlsplit(url)
    host = (parsed.hostname or "").casefold().removeprefix("www.")
    path = parsed.path.casefold()
    source_type = str(item.get("source_type") or "").casefold()
    score = 10
    if host.endswith(".gov") or host.endswith(".us"):
        score += 70
    if "linkedin.com/in/" in url.casefold():
        score += 55
    if source_type in {"linkedin_existing", "web_search"}:
        score += 15
    if any(domain in host for domain in ("whitepages.com", "zoominfo.com", "rocketreach.co")):
        score -= 20
    if "linkedin.com/pub/dir" in url.casefold() or (host == "google.com" and path.startswith("/goto")):
        score -= 60
    return score


def _evidence_label(item: dict[str, Any] | None) -> str:
    if not item:
        return ""
    title = " ".join(str(item.get("source_title") or "").split()).strip()
    if title:
        return title[:250]
    host = (urlsplit(str(item.get("source_url") or "")).hostname or "Public source").removeprefix("www.")
    return host


def _friendly_source_type(value: Any) -> str:
    text = str(value or "Public source").replace("_", " ").strip()
    return text.title().replace("Linkedin", "LinkedIn")


def _pick(values: dict[str, Any], *aliases: str) -> Any:
    lookup = {str(key).strip().casefold(): value for key, value in values.items()}
    for alias in aliases:
        if alias.casefold() in lookup:
            return lookup[alias.casefold()]
    return None


def _unique_headers(headers: list[str]) -> list[str]:
    seen: Counter[str] = Counter()
    output: list[str] = []
    for header in headers:
        label = " ".join(str(header or "Column").split()) or "Column"
        seen[label.casefold()] += 1
        output.append(label if seen[label.casefold()] == 1 else f"{label} ({seen[label.casefold()]})")
    return output


def _status_label(status: str | None) -> str:
    return {
        "completed": "Completed",
        "manual_review": "Review Required",
        "failed": "Failed",
        "processing": "Processing",
        "pending": "Pending",
    }.get(str(status), "Pending")


def _safe_value(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value
