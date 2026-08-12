from __future__ import annotations

import csv
import re
import io
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from securedesk.models import ContactRow


ALIASES = {
    "name": {
        "name",
        "contact",
        "contact name",
        "full name",
        "person",
        "person name",
    },
    "organization": {
        "organization",
        "organisation",
        "organization name",
        "organisation name",
        "company",
        "company name",
        "employer",
        "agency",
        "municipality",
        "account",
        "account name",
    },
    "title": {
        "title",
        "job title",
        "position",
        "role",
        "contact title",
        "position title",
    },
    "email": {
        "email",
        "email address",
        "email id",
        "e-mail",
        "work email",
        "work e-mail",
        "contact email",
    },
    "location": {
        "location",
        "market",
        "region",
        "mailing address",
        "address",
        "city state",
        "contact location",
        "territory",
    },
    "city": {"city", "municipality name", "contact city", "organization city"},
    "state": {"state", "state code", "province", "contact state"},
    "linkedin_url": {
        "linkedin",
        "linkedin url",
        "linkedin profile",
        "linkedin profile url",
        "profile url",
    },
    "linkedin_status": {"linkedin lookup status"},
    "linkedin_note": {"linkedin validation note"},
    "linkedin_confidence": {"linkedin match confidence"},
    "linkedin_title": {"linkedin title", "linkedin job title"},
    "linkedin_snippet": {"linkedin snippet", "linkedin description"},
    "flag": {"flag", "flagged", "manual review"},
    "flag_reason": {"flag reason", "review reason"},
    "organization_website": {
        "organization website",
        "official website",
        "website",
        "company website",
        "domain",
    },
}


def normalized_header(value: Any) -> str:
    return " ".join(str(value or "").strip().casefold().replace("_", " ").split())


def _unique_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = " ".join(str(value or "").split()).strip() or f"Column {index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base} ({counts[base]})")
    return headers


def _column_mapping(headers: list[str]) -> dict[str, int]:
    reverse = {
        alias: canonical
        for canonical, aliases in ALIASES.items()
        for alias in aliases
    }
    mapping: dict[str, int] = {}
    for index, header in enumerate(headers):
        canonical = reverse.get(normalized_header(header))
        if canonical and canonical not in mapping:
            mapping[canonical] = index
    return mapping


def detected_columns(headers: list[str]) -> dict[str, str]:
    return {
        canonical: headers[index]
        for canonical, index in _column_mapping(headers).items()
        if index < len(headers)
    }


def _require_columns(headers: list[str]) -> dict[str, int]:
    mapping = _column_mapping(headers)
    missing = [field for field in ("name", "organization") if field not in mapping]
    if missing:
        raise ValueError(
            "Could not find the required column(s): "
            + ", ".join(missing)
            + ". Rename them to Name and Organization."
        )
    return mapping


def _make_row(
    *,
    row_id: int,
    source_row: int,
    headers: list[str],
    values: list[Any],
    mapping: dict[str, int],
) -> ContactRow:
    padded = [*values, *([None] * max(0, len(headers) - len(values)))]
    original = {
        header: padded[index] if index < len(padded) else None
        for index, header in enumerate(headers)
    }
    def optional_value(field: str) -> Any:
        index = mapping.get(field)
        return padded[index] if index is not None and index < len(padded) else None

    existing_linkedin = str(optional_value("linkedin_url") or "").strip()
    valid_existing_linkedin = (
        existing_linkedin
        if existing_linkedin.casefold().startswith((
            "https://linkedin.com/in/",
            "http://linkedin.com/in/",
            "https://www.linkedin.com/in/",
            "http://www.linkedin.com/in/",
        ))
        else None
    )
    raw_confidence = optional_value("linkedin_confidence")
    try:
        existing_confidence = int(float(raw_confidence or 0))
    except (TypeError, ValueError):
        existing_confidence = 0

    missing_identity_fields = [
        label
        for field, label in (
            ("name", "Name"),
            ("organization", "Organization"),
            ("title", "Title"),
            ("email", "Email"),
        )
        if not str(optional_value(field) or "").strip()
    ]
    raw_prior_flag = str(optional_value("flag") or "").strip().casefold() in {
        "1",
        "true",
        "yes",
        "y",
        "flagged",
    }
    incomplete_reason = (
        f"Incomplete information: missing {', '.join(missing_identity_fields)}."
        if missing_identity_fields
        else None
    )
    raw_prior_flag_reason = str(optional_value("flag_reason") or "").strip() or None
    prior_reasons = [
        part.strip()
        for part in re.split(r"\s*;\s*", raw_prior_flag_reason or "")
        if part.strip()
        and part.strip().casefold()
        not in {"linkedin not found", "user linkedin not found"}
    ]
    prior_flag_reason = "; ".join(prior_reasons) or None
    linkedin_miss_was_only_flag = bool(raw_prior_flag_reason) and not prior_flag_reason
    prior_flag = raw_prior_flag and not linkedin_miss_was_only_flag

    return ContactRow(
        row_id=row_id,
        source_row=source_row,
        name=padded[mapping["name"]] if mapping["name"] < len(padded) else None,
        organization=(
            padded[mapping["organization"]]
            if mapping["organization"] < len(padded)
            else None
        ),
        title=optional_value("title"),
        email=optional_value("email"),
        location=(
            optional_value("location")
            or " ".join(
                str(value).strip()
                for value in (optional_value("city"), optional_value("state"))
                if value
            )
            or None
        ),
        city=optional_value("city"),
        state=optional_value("state"),
        linkedin_title=optional_value("linkedin_title"),
        linkedin_snippet=optional_value("linkedin_snippet"),
        organization_website=optional_value("organization_website"),
        initial_linkedin_url=valid_existing_linkedin,
        initial_lookup_status="found" if valid_existing_linkedin else "pending",
        initial_validation_note=(
            optional_value("linkedin_note")
            if valid_existing_linkedin
            else "Previous result will be retried."
            if optional_value("linkedin_status") or existing_linkedin
            else None
        ),
        initial_match_confidence=(
            max(1, min(existing_confidence or 100, 100))
            if valid_existing_linkedin
            else 0
        ),
        initial_flagged=bool(missing_identity_fields or prior_flag),
        initial_flag_reason="; ".join(
            reason for reason in (incomplete_reason, prior_flag_reason) if reason
        ) or None,
        original_data=original,
    )


def read_csv_file(path: Path) -> tuple[list[str], list[ContactRow]]:
    raw = path.read_bytes()
    text: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("The CSV encoding is not supported. Save it as UTF-8 CSV.")

    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        raw_headers = next(reader)
    except StopIteration as error:
        raise ValueError("The CSV file is empty") from error
    headers = _unique_headers(raw_headers)
    mapping = _require_columns(headers)
    rows: list[ContactRow] = []
    for source_row, values in enumerate(reader, start=2):
        if not any(str(value or "").strip() for value in values):
            continue
        rows.append(
            _make_row(
                row_id=len(rows) + 1,
                source_row=source_row,
                headers=headers,
                values=list(values),
                mapping=mapping,
            )
        )
    if not rows:
        raise ValueError("The CSV file contains no contact rows")
    return headers, rows


def read_xlsx_file(
    path: Path,
    *,
    header_scan_rows: int = 25,
) -> tuple[list[str], list[ContactRow]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        last_error: ValueError | None = None
        for worksheet in workbook.worksheets:
            for header_row, raw_values in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=min(header_scan_rows, worksheet.max_row or header_scan_rows),
                    values_only=True,
                ),
                start=1,
            ):
                headers = _unique_headers(list(raw_values))
                try:
                    mapping = _require_columns(headers)
                except ValueError as error:
                    last_error = error
                    continue
                rows: list[ContactRow] = []
                for source_row, values in enumerate(
                    worksheet.iter_rows(min_row=header_row + 1, values_only=True),
                    start=header_row + 1,
                ):
                    values_list = list(values)
                    if not any(str(value or "").strip() for value in values_list):
                        continue
                    rows.append(
                        _make_row(
                            row_id=len(rows) + 1,
                            source_row=source_row,
                            headers=headers,
                            values=values_list,
                            mapping=mapping,
                        )
                    )
                if not rows:
                    raise ValueError("The Excel sheet contains no contact rows")
                return headers, rows
        raise last_error or ValueError(
            "Could not find Name and Organization columns in the Excel workbook"
        )
    finally:
        workbook.close()


def read_contact_file(
    path: Path,
    *,
    header_scan_rows: int = 25,
) -> tuple[list[str], list[ContactRow]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return read_csv_file(path)
    if suffix == ".xlsx":
        return read_xlsx_file(path, header_scan_rows=header_scan_rows)
    raise ValueError("Only .csv and .xlsx files are supported")
